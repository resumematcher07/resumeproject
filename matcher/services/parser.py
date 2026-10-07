"""Read an uploaded resume end to end and split it into categories."""
import io
import re
import unicodedata
from datetime import date

from .skills_data import find_certs, find_skills, flat_skills

# ---------- text extraction ----------


def extract_text(uploaded_file, filename):
    ext = filename.rsplit(".", 1)[-1].lower()
    data = uploaded_file.read()
    if ext == "pdf":
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        return "\n".join((page.extract_text() or "") for page in reader.pages)
    if ext == "docx":
        from docx import Document

        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:                       # resumes often hide content in tables
            for row in table.rows:
                for cell in row.cells:
                    parts.extend(p.text for p in cell.paragraphs)
        return "\n".join(parts)
    if ext in ("txt", "md"):
        return data.decode("utf-8", errors="ignore")
    raise ValueError("Unsupported file type. Upload a PDF, DOCX or TXT resume.")


# PDF fonts often drop the "fi" glyph, so these words arrive without it.
_FI_REPAIRS = {
    r"\bCertication(s?)\b": r"Certification\1", r"\bcertication(s?)\b": r"certification\1",
    r"\bnotication(s?)\b": r"notification\1", r"\bprot margin": "profit margin",
    r"\bsignicant": "significant", r"\bspecication": "specification", r"\bqualication": "qualification",
}


def repair_text(text):
    """Undo common PDF text-extraction damage: ligatures, '|' read as 'j', '{' bullets."""
    text = unicodedata.normalize("NFKC", text.replace(" ", " "))
    text = re.sub(r"(?<=[A-Za-z])[\x0c\x0b�](?=[A-Za-z])", "fi", text)   # the PDF "fi" glyph arrives as a control char
    text = text.replace("\x0c", "\n").replace("\x0b", " ")
    text = re.sub(r"(?m)^\s*[{•●▪■◦○·⁃∙]\s*", "• ", text)
    text = re.sub(r"(?<=\S) [jJ] (?=\S)", " | ", text)               # "a j b" is a pipe separator
    for rx, rep in _FI_REPAIRS.items():
        text = re.sub(rx, rep, text)
    return text


# ---------- sectioning ----------

_MORE = r"(?:\s*(?:&|and|/|,)\s*[a-z]+(?:\s+[a-z]+){0,2})?"
HEADERS = {
    "summary": r"(professional |career |personal )?(summary|profile|objective|statement)|about( me)?",
    "skills": r"(technical |key |core |it |soft |professional |hard )?(skills|competenc(?:y|ies))( & tools| and tools)?|technologies|tech stack|tools( & technologies)?",
    "experience": r"(work |professional |employment |relevant |industry |internship )?(experience|history)s?|employment|work history|career history|internships?",
    "education": r"education(al)?( background| qualifications?)?|academics?|academic (background|qualifications?)|qualifications?",
    "certifications": r"(licen[cs]es? (&|and) )?certifications?|certificates?|licen[cs]es?|courses|training|professional development",
    "projects": r"(personal |academic |key |selected |notable |mini )?projects?",
    "achievements": r"achievements?|awards?( & honou?rs)?|honou?rs|accomplishments?",
    "languages": r"languages?",
    "activities": r"(community|volunteer|leadership|extra[- ]?curricular|social)[a-z &/]*|activities|positions of responsibility",
}
_HEADER_RX = {k: re.compile(rf"^(?:{v}){_MORE}$", re.I) for k, v in HEADERS.items()}

BULLET = re.compile(r"^\s*(?:[•●▪■◦○\-–—*·►✓➢{]|\d+[.)])\s*")


def _is_header(line):
    label = re.sub(r"[^A-Za-z&/, ]", "", line.strip()).strip()
    return bool(label) and len(line.split()) <= 8 and any(rx.match(label) for rx in _HEADER_RX.values())


def split_sections(text):
    sections = {"header": []}
    current = "header"
    for raw in text.splitlines():
        line = raw.strip()
        if not line:
            continue
        label = re.sub(r"[^A-Za-z&/, ]", "", line).strip()
        matched = None
        if 3 <= len(label) <= 60 and len(line.split()) <= 8 and not BULLET.match(line) and not line.endswith("."):
            for key, rx in _HEADER_RX.items():
                if rx.match(label):
                    matched = key
                    break
        if matched:
            current = matched
            sections.setdefault(current, [])
            continue
        sections[current].append(line)
    return sections


# ---------- dates & experience ----------

MONTH = r"(?:jan(?:uary)?|feb(?:ruary)?|mar(?:ch)?|apr(?:il)?|may|june?|july?|aug(?:ust)?|sept?(?:ember)?|oct(?:ober)?|nov(?:ember)?|dec(?:ember)?)"
YEAR = r"(?:19|20)\d{2}"
MONTHS = {m: i for i, m in enumerate(
    ["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"], 1)}
RANGE = re.compile(
    rf"(?<![A-Za-z])(?:(?P<m1>{MONTH})\.?,?\s+)?(?P<y1>{YEAR})\s*(?:-|–|—|to|until|till)\s*"
    rf"(?:(?:(?P<m2>{MONTH})\.?,?\s+)?(?P<y2>{YEAR})|(?P<now>present|current|currently|now|till date|to date|ongoing|today|date))",
    re.I,
)
SINGLE = re.compile(rf"(?<![A-Za-z])(?P<m>{MONTH})\.?,?\s+(?P<y>{YEAR})\b", re.I)
NUMERIC_DATE = re.compile(r"\b(0?[1-9]|1[0-2])\s*[/.]\s*(" + YEAR + r")\b")
_MON_NAMES = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]

STATED_PATTERNS = [
    re.compile(r"(?P<n>\d{1,2}(?:\.\d)?)\s*\+?\s*(?P<u>years?|yrs?|months?)\s+(?:of\s+)?(?:[\w-]+\s+){0,3}?experience", re.I),
    re.compile(r"experience\s*(?:of|:|-)?\s*(?:over|around|about|nearly)?\s*(?P<n>\d{1,2}(?:\.\d)?)\s*\+?\s*(?P<u>years?|yrs?|months?)", re.I),
]


def _month_index(tok):
    return MONTHS.get((tok or "")[:3].lower())


def _normalise_dates(text):
    return NUMERIC_DATE.sub(lambda m: f"{_MON_NAMES[int(m.group(1)) - 1]} {m.group(2)}", text)


def has_date(text):
    t = _normalise_dates(text)
    return bool(RANGE.search(t) or SINGLE.search(t))


def month_spans(lines, today=None):
    """Inclusive (first_month .. last_month) spans, one per dated role, as month indexes.

    - 'Aug 2025 - Jan 2026'  -> Aug..Jan inclusive = 6 months
    - 'Apr 2026 - Present'   -> Apr..current month
    - 'May 2023' (one month) -> 1 month
    - '2021 - 2023' (years)  -> Jan 2021 .. Dec 2023
    """
    today = today or date.today()
    now_idx = today.year * 12 + today.month
    spans = []
    for raw in lines:
        line = _normalise_dates(raw)
        for m in RANGE.finditer(line):
            start = int(m["y1"]) * 12 + (_month_index(m["m1"]) or 1)
            if m["now"]:
                end = now_idx
            else:
                end = int(m["y2"]) * 12 + (_month_index(m["m2"]) or 12)
            end = min(end, now_idx)                        # a future end date cannot add experience
            if start <= end and end - start + 1 <= 12 * 45:
                spans.append((start, end))
        if not BULLET.match(raw) and len(raw) <= 140:      # a lone "May 2023" on a role line = that month
            for m in SINGLE.finditer(RANGE.sub(" ", line)):
                idx = int(m["y"]) * 12 + _month_index(m["m"])
                if idx <= now_idx:
                    spans.append((idx, idx))
    return spans


def months_worked(spans):
    """Total distinct months across spans (overlapping roles are counted once)."""
    merged = []
    for a, b in sorted(spans):
        if merged and a <= merged[-1][1]:
            merged[-1][1] = max(merged[-1][1], b)
        else:
            merged.append([a, b])
    return sum(b - a + 1 for a, b in merged)


def months_span(spans):
    return (max(b for _, b in spans) - min(a for a, _ in spans) + 1) if spans else 0


def stated_experience_months(text):
    """Experience the candidate wrote down themselves ('3 years of experience') - this wins over calculation."""
    best = 0
    for rx in STATED_PATTERNS:
        for m in rx.finditer(text):
            n = float(m["n"])
            months = round(n) if m["u"].lower().startswith("month") else round(n * 12)
            if 0 < months <= 12 * 45:
                best = max(best, months)
    return best


def experience_label(months):
    y, m = divmod(int(months), 12)
    if not months:
        return "No experience dates"
    if not y:
        return f"{m} mo"
    return f"{y} yr" + (f" {m} mo" if m else "")


def compute_experience(experience_lines, summary_text, today=None):
    from django.conf import settings

    spans = month_spans(experience_lines, today)
    worked, span = months_worked(spans), months_span(spans)
    calculated = span if getattr(settings, "EXPERIENCE_METHOD", "worked") == "span" else worked
    stated = stated_experience_months(summary_text)
    months, source = (stated, "stated in resume") if stated else (calculated, "calculated from dates")
    return {
        "experience_months": months,
        "years_experience": round(months / 12, 1),
        "experience_label": experience_label(months),
        "experience_source": source if months else "none found",
        "experience_worked_months": worked,
        "experience_span_months": span,
    }


# ---------- titles ----------

TITLE_WORDS = (
    r"engineer|developer|programmer|architect|analyst|scientist|manager|director|designer|consultant|"
    r"administrator|specialist|lead|head|officer|executive|associate|intern|trainee|tester|qa|sdet|devops|sre|"
    r"accountant|auditor|coordinator|recruiter|writer|marketer|strategist|technician|supervisor|assistant|"
    r"representative|advisor|owner|founder|professor|teacher|nurse|researcher|statistician|developer"
)
TITLE_RX = re.compile(rf"\b(?:{TITLE_WORDS})\b", re.I)
_TITLE_WORD_FULL = re.compile(rf"(?:{TITLE_WORDS})", re.I)
_LEVEL_WORD = re.compile(r"(?:entry[- ]?level|junior|jr\.?|senior|sr\.?|lead|principal|staff|associate|fresher|aspiring|mid[- ]?level)", re.I)
_NOT_PART_OF_TITLE = {"a", "an", "the", "as", "with", "and", "of", "in", "for", "to", "at", "by", "on", "who", "that",
                      "is", "am", "seeking", "looking", "skilled", "experienced", "motivated", "dedicated", "result-oriented",
                      "highly", "passionate", "detail-oriented", "driven", "i", "my", "our", "from", "like"}


def title_from_sentence(sentence):
    words = re.findall(r"[A-Za-z][A-Za-z+#&/.'-]*", sentence)
    for i, w in enumerate(words[:12]):
        if not _TITLE_WORD_FULL.fullmatch(w.rstrip(".")):
            continue
        j = i
        while j > 0 and i - j < 4:
            prev = words[j - 1]
            if prev.lower() in _NOT_PART_OF_TITLE:
                break
            if prev[0].isupper() or _LEVEL_WORD.fullmatch(prev):
                j -= 1
            else:
                break
        return " ".join(words[j:i + 1]).replace(" & ", " & ")
    return None


def title_from_summary(summary_text):
    """The job a candidate says they are is almost always in the first sentence of the summary."""
    sentences = re.split(r"(?<=[.!?])\s+", summary_text.strip())
    for s in sentences[:2]:
        t = title_from_sentence(s)
        if t:
            return t
    return None


_DATE_STRIP = re.compile(rf"(?<![A-Za-z])(?:{MONTH}\.?,?\s+)?{YEAR}\b(?:\s*(?:-|–|—|to)\s*(?:(?:{MONTH}\.?,?\s+)?{YEAR}|present|current|now|till date))?", re.I)


def strip_dates(text):
    return re.sub(r"\s{2,}", " ", _DATE_STRIP.sub(" ", _normalise_dates(text))).strip(" -–—|,:")


def _clean_title(line):
    line = re.sub(r"\(.*?\)", "", strip_dates(line))
    parts = re.split(r"\s*[|•·–—]\s*|\s+-\s+|\s+(?:at|@)\s+|,\s+|\s{3,}", line)
    for p in parts:
        p = p.strip(" -–—|,:")
        if p and TITLE_RX.search(p) and 2 <= len(p) <= 60 and len(p.split()) <= 7:
            return p
    return None


def find_titles(exp_lines, header_lines):
    titles = []
    for line in exp_lines:
        if BULLET.match(line) or len(line) > 130 or line[:1].islower():
            continue
        t = _clean_title(line)
        if t and t.lower() not in [x.lower() for x in titles]:
            titles.append(t)
    headline = None
    for line in header_lines[:6]:
        if EMAIL.search(line) or PHONE.search(line) or len(line) > 70:
            continue
        t = _clean_title(line)
        if t:
            headline = t
            break
    return titles, headline


# ---------- entries (role / project blocks) ----------


def group_entries(lines):
    """Turn raw experience/project lines into [{'head': [...], 'bullets': [...]}].

    PDF text wraps long bullets onto extra lines; those continuation lines are glued back onto their bullet.
    """
    entries, cur = [], None
    for raw in lines:
        line = raw.strip()
        if not line:
            continue
        is_bullet = bool(BULLET.match(line))
        text = BULLET.sub("", line).strip()
        if not text:
            continue
        if is_bullet:
            if cur is None:
                cur = {"head": [], "bullets": []}
                entries.append(cur)
            cur["bullets"].append(text)
            continue
        if cur is None:
            cur = {"head": [text], "bullets": []}
            entries.append(cur)
        elif cur["bullets"]:
            heading_like = has_date(text) or ("|" in text and len(text) <= 140)
            ended = bool(re.search(r"[.!?)%]$", cur["bullets"][-1]))
            if text[0].islower():
                cur["bullets"][-1] += " " + text
            elif heading_like or (ended and len(text) <= 90):
                cur = {"head": [text], "bullets": []}
                entries.append(cur)
            else:
                cur["bullets"][-1] += " " + text
        else:
            if has_date(text) and any(has_date(h) for h in cur["head"]):
                cur = {"head": [text], "bullets": []}
                entries.append(cur)
            else:
                cur["head"].append(text)
    return entries


def split_date(text):
    """('Data Scientist | Aug 2025 - Jan 2026') -> ('Data Scientist', 'Aug 2025 - Jan 2026')."""
    t = _normalise_dates(text)
    m = RANGE.search(t) or SINGLE.search(t)
    if not m:
        return text.strip(" |"), ""
    return strip_dates(t), re.sub(r"\s+", " ", m.group(0)).strip()


# ---------- field helpers ----------

EMAIL = re.compile(r"[\w.+-]+@[\w-]+\.[\w.-]+")
PHONE = re.compile(r"(?<!\d)(\+?\d[\d\s().-]{8,16}\d)(?!\d)")
LINK = re.compile(r"(?:https?://)?(?:www\.)?(?:linkedin\.com/in|github\.com|gitlab\.com|behance\.net|[\w-]+\.(?:dev|io|me))/[\w\-./%]+", re.I)
DEGREE = re.compile(
    r"\b(b\.?\s?tech|m\.?\s?tech|b\.?e\b|m\.?e\b|b\.?sc|m\.?sc|b\.?s\b|m\.?s\b|bca|mca|mba|ph\.?d|bachelor|master|diploma|associate degree|b\.?com|m\.?com|b\.?a\b|m\.?a\b)",
    re.I,
)
INLINE_CERTS = re.compile(r"(?i)\b(?:certifications?|certificates?|courses)(?:\s+(?:and|&)\s+(?:courses|certifications?))?\s*:\s*(.+)")


def guess_name(header_lines):
    for line in header_lines[:5]:
        words = line.split()
        if 2 <= len(words) <= 4 and all(re.fullmatch(r"[A-Za-z.'\-]+", w) for w in words) and not TITLE_RX.search(line):
            return line.title() if line.isupper() else line
    return "Candidate"


def unbullet(lines):
    return [BULLET.sub("", l).strip() for l in lines if l.strip()]


def split_items(lines, seps=r"\s*\|\s*|\s*;\s*"):
    joined, buf = [], ""
    for l in unbullet(lines):                       # rejoin "(HP" + "Life)" style wrapped brackets
        buf = f"{buf} {l}".strip()
        if buf.count("(") <= buf.count(")"):
            joined.append(buf)
            buf = ""
    if buf:
        joined.append(buf)
    out = []
    for l in joined:
        out.extend(p.strip() for p in re.split(seps, l) if p.strip())
    return out


# ---------- main entry ----------


def parse_resume(text):
    text = repair_text(text)
    sections = split_sections(text)
    header = sections.get("header", [])
    exp = sections.get("experience", [])
    summary_text = " ".join(unbullet(sections.get("summary", [])))

    body = [l for k, ls in sections.items() if k != "header" for l in ls]      # name/contact initials are not skills
    skills_found = find_skills("\n".join(body) if body else text)
    flat = flat_skills(skills_found)

    titles, headline = find_titles(exp, header)
    latest_role = titles[0] if titles else None
    summary_title = title_from_summary(summary_text)
    primary = summary_title or headline or latest_role              # the job the candidate says they are
    all_titles = list(dict.fromkeys(t for t in [primary, headline, *titles] if t))

    # experience is read ONLY from the experience box; without one, skip education/projects/volunteering
    if exp:
        exp_basis = exp
    else:
        skip = {"education", "projects", "certifications", "activities", "achievements", "languages", "summary"}
        exp_basis = [l for k, ls in sections.items() if k not in skip for l in ls]
    exp_info = compute_experience(exp_basis, " ".join(header[:6]) + " " + summary_text)

    cert_lines = split_items(sections.get("certifications", []))
    all_lines = text.splitlines()
    for i, line in enumerate(all_lines):                           # "Certifications and courses: A; B; C"
        m = INLINE_CERTS.search(line)
        if m:
            chunk = m.group(1).strip()
            while not chunk.endswith(".") and i + 1 < len(all_lines) and len(chunk) < 400 and not _is_header(all_lines[i + 1]):
                i += 1
                chunk += " " + all_lines[i].strip()
            cert_lines.extend(p.strip(" .") for p in re.split(r"\s*;\s*", chunk) if p.strip(" ."))
    cert_lines = list(dict.fromkeys(cert_lines))
    certs_known = find_certs(text)
    edu_lines = unbullet(sections.get("education", []))
    if not edu_lines:
        edu_lines = [l for l in text.splitlines() if DEGREE.search(l)][:4]

    phone = next((m.group(1).strip() for m in PHONE.finditer("\n".join(header[:8]) or text) if len(re.sub(r"\D", "", m.group(1))) >= 9), "")
    em = EMAIL.search(text)

    return {
        "name": guess_name(header),
        "email": em.group(0) if em else "",
        "phone": phone,
        "links": sorted({m.group(0) for m in LINK.finditer(text)})[:4],
        "summary": summary_text[:900],
        "skills": skills_found,
        "skills_flat": flat,
        "skills_section": unbullet(sections.get("skills", [])),
        **exp_info,
        "primary_title": primary,
        "title_source": "summary" if summary_title else ("headline" if headline else ("experience" if latest_role else "")),
        "latest_title": latest_role or primary,
        "titles": all_titles[:6],
        "experience": exp,
        "education": edu_lines[:10],
        "certifications": cert_lines[:20],
        "certs_detected": certs_known,
        "projects": sections.get("projects", [])[:60],
        "achievements": unbullet(sections.get("achievements", []))[:10],
        "activities": sections.get("activities", [])[:12],
        "languages": unbullet(sections.get("languages", []))[:6],
        "sections_found": [k for k in sections if k != "header" and sections[k]],
        "word_count": len(text.split()),
    }
