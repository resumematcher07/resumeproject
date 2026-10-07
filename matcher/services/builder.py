"""Build a fresh, role-tailored resume (.docx) from the parsed original.

Only facts the candidate already has - or explicitly confirms on the results page - are written into
the new resume. Skills the candidate does not have are never invented; they go to the learning plan.
"""
import io
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from .parser import group_entries, split_date

ACCENT = RGBColor(0x4F, 0x46, 0xE5)


def _heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(text.upper())
    r.bold = True
    r.font.size = Pt(10.5)
    r.font.color.rgb = ACCENT
    pPr = p._p.get_or_add_pPr()                               # thin rule under each heading
    bdr = pPr.makeelement(qn("w:pBdr"), {})
    bottom = bdr.makeelement(qn("w:bottom"), {qn("w:val"): "single", qn("w:sz"): "4", qn("w:space"): "1", qn("w:color"): "C7D2FE"})
    bdr.append(bottom)
    pPr.append(bdr)


def _line(doc, text, bullet=False, bold=False):
    p = doc.add_paragraph(style="List Bullet" if bullet else None)
    p.paragraph_format.space_after = Pt(1)
    r = p.add_run(text)
    r.bold = bold
    r.font.size = Pt(10)
    return p


def build_resume(parsed, job, evaluation, extra_skills=(), pursuing_certs=()):
    extra_skills = [s for s in extra_skills if s in set(evaluation["missing_required"] + evaluation["missing_preferred"])]
    pursuing = [c for c in pursuing_certs if c in {x["name"] for x in evaluation["missing_certs"]}]

    doc = Document()
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Pt(54)
    sec.top_margin = sec.bottom_margin = Pt(46)
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(parsed.get("name") or "Candidate")
    r.bold = True
    r.font.size = Pt(22)
    contact = " | ".join(x for x in [parsed.get("email"), parsed.get("phone"), *parsed.get("links", [])] if x)
    if contact:
        _line(doc, contact)
    _line(doc, f"Targeting: {job['title']} at {job['company']}", bold=True)

    # --- tailored summary (built only from true / confirmed data) ---
    top = (evaluation["matched"] + extra_skills)[:6]
    years = parsed.get("years_experience") or 0
    role = parsed.get("primary_title") or parsed.get("latest_title") or "Professional"
    summary = f"{role}" + (f" with {parsed.get('experience_label')} of experience" if parsed.get("experience_months") else "")
    if top:
        summary += f" skilled in {', '.join(top[:-1]) + (' and ' if len(top) > 1 else '') + top[-1]}"
    summary += f". Seeking to contribute to {job['company']} as {job['title']}."
    if parsed.get("summary"):
        summary += " " + parsed["summary"]
    _heading(doc, "Professional Summary")
    _line(doc, summary)

    # --- skills: role-relevant first ---
    relevant = list(dict.fromkeys(evaluation["matched"] + extra_skills))
    others = [s for s in parsed.get("skills_flat", []) if s not in relevant]
    _heading(doc, "Skills")
    if relevant:
        _line(doc, "Key skills for this role: " + ", ".join(relevant), bold=True)
    if others:
        _line(doc, "Additional skills: " + ", ".join(others))

    def entries(title, lines):
        if not lines:
            return
        _heading(doc, title)
        for e in group_entries(lines):
            head = []
            for h in e["head"]:
                text, d = split_date(h)
                head.append(text + (f"  ({d})" if d else ""))
            if head:
                _line(doc, " | ".join(head), bold=True)
            for b in e["bullets"]:
                _line(doc, b, bullet=True)

    entries("Professional Experience", parsed.get("experience"))
    entries("Projects", parsed.get("projects"))

    if parsed.get("education"):
        _heading(doc, "Education")
        for l in parsed["education"]:
            _line(doc, l)

    certs = parsed.get("certifications", [])
    if certs or pursuing:
        _heading(doc, "Certifications")
        for c in certs:
            _line(doc, c, bullet=True)
        for c in pursuing:
            _line(doc, f"{c} (in progress)", bullet=True)

    if parsed.get("achievements"):
        _heading(doc, "Achievements")
        for l in parsed["achievements"]:
            _line(doc, l, bullet=True)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def safe_filename(name, job):
    base = re.sub(r"[^A-Za-z0-9]+", "_", f"{name}_{job['company']}_{job['title']}").strip("_")
    return f"{base[:90]}.docx"
